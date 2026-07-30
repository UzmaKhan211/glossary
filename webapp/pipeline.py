"""
Glue between the existing command-line scripts (scripts/extract_data.py,
scripts/pdf_glossary_parser.py, scripts/generate.py) and the Flask app in
app.py. Nothing in here is Flask-specific -- app.py just calls these
functions and renders their results.

Data lives in the same repo-relative paths the CLI scripts already use
(data/vocab.json, data/unit_vocab.json, output/*.pdf) so the web interface
is a front end for the real pipeline, not a separate copy of it. The only
new persistent state is the uploaded PDFs themselves and their parsed
cache, under data/unit_pdfs/ and data/unit_glossary_parsed/.
"""
import json
import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import extract_data  # noqa: E402
from pdf_glossary_parser import GlossaryEntry, parse_unit_pdf  # noqa: E402
from grade_fonts import grade_css  # noqa: E402

REPO_ROOT = Path(__file__).resolve().parent.parent
DATA_DIR = REPO_ROOT / "data"
OUTPUT_DIR = REPO_ROOT / "output"
UNIT_PDFS_DIR = DATA_DIR / "unit_pdfs"
UNIT_PARSED_DIR = DATA_DIR / "unit_glossary_parsed"
VOCAB_PATH = DATA_DIR / "vocab.json"
UNIT_VOCAB_PATH = DATA_DIR / "unit_vocab.json"

# The grades extract_data.py's fixed column layout currently understands.
# This ISN'T a hard cap on what the app can show -- known_grades() below
# also picks up whatever grades actually exist in data/vocab.json, so a
# grades 6-8 spreadsheet (once extract_data.py -- or a variant of it -- can
# read that sheet's own column layout) shows up here automatically without
# any other code change. Uploading that sheet through the *current*
# extractor won't work until its column layout is wired in, since
# extract_data.GRADE_COLS is specific to the K-5 sheet's fixed columns.
GRADES = ["K", "1", "2", "3", "4", "5"]
GRADE_LABELS = {"K": "Kindergarten", "1": "Grade 1", "2": "Grade 2", "3": "Grade 3", "4": "Grade 4", "5": "Grade 5"}


def label_for_grade(grade):
    return GRADE_LABELS.get(grade, "Kindergarten" if grade == "K" else f"Grade {grade}")


def known_grades():
    """Every grade with word data loaded, in a sensible display order (K-5
    first if present, then any other grade -- e.g. 6, 7, 8 -- numerically)."""
    vocab = load_vocab()
    grades = {g for g in (set(vocab) | set(GRADES)) if vocab.get(g)}
    ordered = [g for g in GRADES if g in grades]
    extra = sorted((g for g in grades if g not in GRADES), key=lambda g: (len(g), g))
    return ordered + extra


# Matches the grade+unit out of Amplify's own PDF filenames, e.g.
# "Grade5-Unit7.pdf", "KG-Unit3.pdf", "Grade 2 - Unit 4.pdf", "Grade7-Unit2.pdf".
FILENAME_RE = re.compile(r"(?:grade\s*(\d+)|k\s*g?)\D*unit\s*(\d+)", re.IGNORECASE)


def grade_and_unit_from_filename(filename):
    m = FILENAME_RE.search(filename)
    if not m:
        return None, None
    grade = m.group(1) if m.group(1) else "K"
    return grade, int(m.group(2))


def load_vocab():
    if VOCAB_PATH.exists():
        return json.loads(VOCAB_PATH.read_text())
    return {g: [] for g in GRADES}


def load_unit_vocab():
    if UNIT_VOCAB_PATH.exists():
        return json.loads(UNIT_VOCAB_PATH.read_text())
    return {}


def save_unit_vocab(unit_vocab):
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    UNIT_VOCAB_PATH.write_text(json.dumps(unit_vocab, indent=2))


def process_master_spreadsheet(xlsx_path):
    """Run the same extraction extract_data.py's CLI does, in-process, and
    merge the result into data/vocab.json. Returns {grade: card_count}.

    This only ever touches the grades extract_data.GRADE_COLS knows how to
    read from this specific spreadsheet's fixed column layout (currently
    K-5) -- any other grade already in data/vocab.json (e.g. 6-8, once a
    matching extractor for that sheet exists) is left untouched rather than
    wiped out, so uploading a second/different spreadsheet is additive."""
    import openpyxl

    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    ws = wb.active
    anchors = extract_data.parse_drawing_anchors(xlsx_path)

    cards_by_grade = {g: [] for g in extract_data.GRADE_COLS}
    for row in range(3, ws.max_row + 1):
        a_raw = ws[f"A{row}"].value
        if not a_raw or not str(a_raw).strip():
            continue
        d_raw = ws[f"D{row}"].value
        word = extract_data.clean_word(a_raw, d_raw, row)

        for grade, (cur_col, img_col, fb_col, new_col) in extract_data.GRADE_COLS.items():
            cols = {"current": cur_col, "image": img_col, "feedback": fb_col, "new_entry": new_col}
            cur_text = ws[f"{cur_col}{row}"].value
            new_text = ws[f"{new_col}{row}"].value
            text_source = "new_entry" if (new_text and str(new_text).strip()) else (
                "current" if (cur_text and str(cur_text).strip()) else None
            )
            raw_text = new_text if text_source == "new_entry" else cur_text
            text = extract_data.clean_definition(raw_text, word) if text_source else ""

            media_name = None
            for role in extract_data.IMAGE_ROLE_PRIORITY:
                key = (row, cols[role])
                if key in anchors:
                    media_name = anchors[key]
                    break

            if not text and not media_name:
                continue

            image_rel_path = None
            if media_name:
                ext = Path(media_name).suffix
                image_rel_path = f"{grade}/{row}_{extract_data.slugify(word)}{ext}"
                extract_data.extract_media(xlsx_path, media_name, extract_data.IMAGES_DIR / image_rel_path)

            cards_by_grade[grade].append({"row": row, "word": word, "text": text, "image": image_rel_path})

    existing = load_vocab()
    existing.update(cards_by_grade)
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    VOCAB_PATH.write_text(json.dumps(existing, indent=2))
    return {g: len(cards) for g, cards in cards_by_grade.items()}


def save_and_parse_unit_pdf(grade, unit, filename, pdf_bytes):
    """Save an uploaded unit PDF and parse it, caching the parsed entries as
    JSON (so re-matching later doesn't need the PDF re-parsed). Returns the
    parsed entry count and any parser warnings."""
    pdf_dir = UNIT_PDFS_DIR / grade
    pdf_dir.mkdir(parents=True, exist_ok=True)
    (pdf_dir / f"unit{unit}.pdf").write_bytes(pdf_bytes)

    parsed = parse_unit_pdf(pdf_bytes)
    parsed_dir = UNIT_PARSED_DIR / grade
    parsed_dir.mkdir(parents=True, exist_ok=True)
    (parsed_dir / f"unit{unit}.json").write_text(json.dumps({
        "source_filename": filename,
        "en": [{"word": e.word, "definition": e.definition} for e in parsed.en],
        "es": [{"word": e.word, "definition": e.definition} for e in parsed.es],
        "warnings": parsed.warnings,
    }, indent=2))
    return len(parsed.en), parsed.warnings


def load_parsed_units(grade):
    """Return {unit_number: {"en": [GlossaryEntry], "es": [GlossaryEntry]}}
    for every unit PDF already uploaded+parsed for a grade."""
    grade_dir = UNIT_PARSED_DIR / grade
    units = {}
    if not grade_dir.is_dir():
        return units
    for path in sorted(grade_dir.glob("unit*.json")):
        m = re.search(r"unit(\d+)\.json$", path.name)
        if not m:
            continue
        data = json.loads(path.read_text())
        units[int(m.group(1))] = {
            "en": [GlossaryEntry(**e) for e in data["en"]],
            "es": [GlossaryEntry(**e) for e in data["es"]],
        }
    return units


def normalize(text):
    text = text.lower().replace("-", " ").replace("(", " ").replace(")", " ")
    return re.sub(r"\s+", " ", text).strip()


def _headword_variants(word):
    base = normalize(word)
    variants = {base}
    if base.endswith("s"):
        variants.add(base[:-1])
    else:
        variants.add(base + "s")
        variants.add(base + "es")
    return variants


def match_grade(grade):
    """Match this grade's cleaned spreadsheet words against every unit PDF
    parsed so far, and write the result into data/unit_vocab.json under
    this grade's key. Returns the list of entries (same shape that gets
    saved) for immediate display."""
    vocab = load_vocab()
    cards = vocab.get(grade, [])
    units = load_parsed_units(grade)

    # unit -> {normalized headword: index into that unit's en/es lists}
    unit_index = {}
    for unit_num, parsed in units.items():
        index = {}
        for i, entry in enumerate(parsed["en"]):
            for variant in _headword_variants(entry.word):
                index.setdefault(variant, i)
        unit_index[unit_num] = index

    entries = []
    for card in cards:
        candidates = sorted(
            unit_num for unit_num, index in unit_index.items()
            if _headword_variants(card["word"]) & set(index)
        )
        primary_unit = candidates[0] if candidates else None
        spanish_word, spanish_text = "", ""
        if primary_unit is not None:
            index = unit_index[primary_unit]
            variant_hit = next(iter(_headword_variants(card["word"]) & set(index)))
            i = index[variant_hit]
            spanish_word = units[primary_unit]["es"][i].word
            spanish_text = units[primary_unit]["es"][i].definition
        entries.append({
            "word": card["word"],
            "text": card.get("text") or "",
            "image": card.get("image"),
            "primary_unit": primary_unit,
            "all_units": candidates,
            "spanish_word": spanish_word,
            "spanish_text": spanish_text,
        })

    unit_vocab = load_unit_vocab()
    unit_vocab[grade] = entries
    save_unit_vocab(unit_vocab)
    return entries


def uploaded_units_summary(grade):
    """{unit_number: entry_count} for whatever's been uploaded+parsed so far."""
    return {unit: len(parsed["en"]) for unit, parsed in load_parsed_units(grade).items()}


def render_card_pdf(grades, language, out_path, year, print_code=""):
    """Render a standalone card deck PDF for the given grades, in English
    (language="en") or Spanish (language="es"), using the same template/CSS
    the original generator uses -- just swapping which fields feed the
    card's headword/definition text."""
    import datetime
    from jinja2 import Environment, FileSystemLoader
    from playwright.sync_api import sync_playwright
    import os

    templates_dir = REPO_ROOT / "templates"
    images_dir = REPO_ROOT / "assets" / "images"
    unit_vocab = load_unit_vocab()

    cards = []
    for grade in grades:
        for entry in unit_vocab.get(grade, []):
            if language == "es":
                word, text = entry.get("spanish_word") or "", entry.get("spanish_text") or ""
                if not word:
                    continue  # no Spanish translation available for this word -- skip rather than print blank
            else:
                word, text = entry["word"], entry.get("text") or ""
            image_uri = None
            if entry.get("image"):
                image_path = images_dir / entry["image"]
                if image_path.exists():
                    image_uri = image_path.resolve().as_uri()
            cards.append({"grade": grade, "word": word, "text": text, "image_uri": image_uri})

    if not cards:
        return 0

    env = Environment(loader=FileSystemLoader(str(templates_dir)))
    template = env.get_template("deck.html.j2")
    pages = [(cards[i], cards[i + 1] if i + 1 < len(cards) else None) for i in range(0, len(cards), 2)]
    html = template.render(pages=pages, grade_css=grade_css(), year=year, print_code=print_code, lang=language)

    build_html_path = templates_dir / f"_build_{language}.html"
    build_html_path.write_text(html)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    browsers_path = os.environ.get("PLAYWRIGHT_BROWSERS_PATH")
    executable_path = None
    if browsers_path:
        candidates = sorted(Path(browsers_path).glob("chromium-*/chrome-linux/chrome"))
        if candidates:
            executable_path = str(candidates[-1])

    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=executable_path)
        page = browser.new_page()
        page.goto(build_html_path.resolve().as_uri())
        page.evaluate(_AUTOFIT_JS, 0.62)
        page.pdf(
            path=str(out_path),
            width="11in",
            height="8.5in",
            print_background=True,
            margin={"top": "0", "bottom": "0", "left": "0", "right": "0"},
        )
        browser.close()
    build_html_path.unlink()
    return len(cards)


_AUTOFIT_JS = """
(minRatio) => {
  function shrinkToFit(card, selector, floorRatio) {
    const el = card.querySelector(selector);
    if (!el) return;
    const cs = getComputedStyle(el);
    const originalSize = parseFloat(cs.fontSize);
    const ratio = parseFloat(cs.lineHeight) / originalSize;
    const floor = originalSize * floorRatio;
    let size = originalSize;
    let guard = 0;
    while (card.scrollHeight > card.clientHeight + 1 && size > floor && guard < 60) {
      size -= 0.5;
      el.style.fontSize = size + 'px';
      el.style.lineHeight = (size * ratio) + 'px';
      guard++;
    }
  }
  document.querySelectorAll('.card:not(.empty)').forEach(card => {
    shrinkToFit(card, '.card-text', minRatio);
    if (card.scrollHeight > card.clientHeight + 1) {
      shrinkToFit(card, '.card-header', minRatio);
    }
  });
}
"""
